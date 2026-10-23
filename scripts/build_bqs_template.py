"""Reconstruct the supplied BQS PDF as an editable, data-free Word template.

Run with the bundled Codex Python (python-docx). The PDF, not a prior DOCX,
is the visual authority. Sample vessel data/photos/signatures are never copied.
"""
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin, section.bottom_margin = Mm(25), Mm(18)
section.left_margin, section.right_margin = Mm(14), Mm(14)
section.header_distance, section.footer_distance = Mm(11), Mm(10)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Arial', Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.05
for name, size in [('Title',20),('Heading 1',12),('Heading 2',10.5)]:
    style = doc.styles[name]
    style.font.name, style.font.size = 'Arial', Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)


def prop(parent, tag, **attrs):
    el = OxmlElement('w:'+tag)
    for key, value in attrs.items():
        el.set(qn('w:'+key),str(value))
    parent.append(el)
    return el


def text(value='', style=None, center=False, size=None, bold=False):
    p = doc.add_paragraph(style=style)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(value)
    r.bold = bold or None
    if size:
        r.font.size = Pt(size)
    return p


def manual(p, name, label=None):
    sdt = OxmlElement('w:sdt')
    pr = prop(sdt,'sdtPr')
    prop(pr,'tag',val='bqs_manual_'+name)
    prop(pr,'alias',val=label or name)
    content = prop(sdt,'sdtContent')
    run = prop(content,'r')
    t = prop(run,'t')
    t.text = '—'
    p._p.append(sdt)


def mark(p,name):
    i = len(doc.element.xpath('//w:bookmarkStart'))+1
    prop(p._p,'bookmarkStart',id=i,name=name)
    prop(p._p,'bookmarkEnd',id=i)


headings = [
    ('s1','1. Background',1),('s2',"2. Vessel’s details (Ship’s Particulars)",1),
    ('s3','3. Bunker Quantity Survey',1),('s4','4. Opening soundings',1),
    ('s4_1',"4.1. Vessel’s tanks",2),('s4_2','4.2. Gross volume - m³',2),
    ('s4_3','4.3. Temperature',2),('s4_4','4.4. Specific gravities',2),
    ('s4_5',"4.5. Barge’s tanks",2),('s4_6','4.6. Gross volume - m³',2),
    ('s4_7','4.7. Temperature',2),('s4_8','4.8. Specific gravities',2),
    ('s5','5. Closing soundings',1),('s5_1',"5.1. Vessel’s tanks",2),
    ('s5_2',"5.2. Barge’s tanks",2),('s6','6. Final figures',1),
    ('s7','7. Photographic report',1),('s7_1',"7.1. Vessel’s tanks",2),
    ('s7_2',"7.2. Barge’s tanks",2),('s7_3','7.3. Sampling',2),
    ('s8','8. Attachments',1),('s8_1','8.1. Time log',2),('s8_2','8.2. Sampling record',2),
]


def heading(name, page=False):
    _, title, level = next(x for x in headings if x[0]==name)
    p = text(title,'Heading '+str(level))
    p.paragraph_format.page_break_before = page
    mark(p,'bqs_'+name)
    return p


def table(headers, widths, size=8, borders=True):
    t = doc.add_table(rows=1,cols=len(headers))
    t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,width in zip(t.columns,widths):
        col.width=Mm(width)
    borders_el=prop(t._tbl.tblPr,'tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        prop(borders_el,edge,val='single' if borders else 'nil',sz=4,color='000000')
    margins=prop(t._tbl.tblPr,'tblCellMar')
    for edge in ['left','right']:
        prop(margins,edge,w=28,type='dxa')
    for cell,width,value in zip(t.rows[0].cells,widths,headers):
        cell.width=Mm(width)
        cell.text=value
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.space_before=Pt(2)
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold=True
                r.font.size=Pt(size)
    prop(t.rows[0]._tr.get_or_add_trPr(),'tblHeader')
    return t


def row(t, values, size=8, center=False):
    cells=t.add_row().cells
    prop(cells[0]._tc.getparent().get_or_add_trPr(),'cantSplit')
    for cell,value in zip(cells,values):
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text=value
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.space_after=Pt(1)
            p.paragraph_format.line_spacing=1
            if center:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size=Pt(size)
    return cells


def tank(prefix):
    p=text('Draft fore: {'+prefix+'_fore} m    Draft aft: {'+prefix+'_aft} m    List: {'+prefix+'_list}    Trim correction: {'+prefix+'_trim_corrected}',size=8)
    p.paragraph_format.keep_with_next=True
    headers=['Tank','Grade','Sounding /\nUllage (m)','Temp\n°C','Total obs\nm³','Water\ndip','Water\nm³','Gross obs\nm³','VCF\n54B','Density\nat 15°C','Std vol\nat 15°C','MT\nIN VAC','MT\nIN AIR']
    t=table(headers,[18,12,17,9,17,10,10,17,13,14,17,14,14],size=7)
    row(t,['{#'+prefix+'_rows}{c0}']+['{c'+str(i)+'}' for i in range(1,12)]+['{c12}{/'+prefix+'_rows}'],size=7,center=True)
    row(t,['TOTAL','','','','','','','{'+prefix+'_gross}','','','{'+prefix+'_standard}','{'+prefix+'_vac}','{'+prefix+'_air}'],size=7,center=True)
    text('Sounding: S = deep sounding; U = ullage. Quantities shown in metric tons.',size=7)


def photo_area(slot, cover=False):
    start=len(doc.element.body)
    if cover:
        text('{#photos_'+slot+'}')
        text('{%image}',center=True)
        text('{/photos_'+slot+'}')
    else:
        gallery=table(['',''],[91,91],borders=False)
        cells=gallery.rows[0].cells
        cells[0].text='{#photos_'+slot+'}'
        cells[0].add_paragraph('{%image}').alignment=WD_ALIGN_PARAGRAPH.CENTER
        cells[1].text='{#right}'
        cells[1].add_paragraph('{%image}').alignment=WD_ALIGN_PARAGRAPH.CENTER
        cells[1].add_paragraph('{/right}')
        cells[1].add_paragraph('{/photos_'+slot+'}')
        prop(gallery.rows[0]._tr.get_or_add_trPr(),'cantSplit')
    text('{^photos_'+slot+'}')
    if cover:
        p=text('Vessel photograph',center=True,size=10)
        p.paragraph_format.space_before=Pt(145)
        p.paragraph_format.space_after=Pt(145)
        p.runs[0].font.color.rgb=RGBColor.from_string('777777')
    else:
        p=text('—',center=True)
    text('{/photos_'+slot+'}')
    nodes=list(doc.element.body)[start-1:-1]
    sdt=OxmlElement('w:sdt')
    pr=prop(sdt,'sdtPr'); prop(pr,'tag',val='bqs_photos_'+slot)
    content=prop(sdt,'sdtContent')
    if nodes:
        nodes[0].addprevious(sdt)
        for node in nodes:
            content.append(node)


# Brand header reuses the repository's logo, never the reference's vessel photos.
header=section.header
ht=header.add_table(rows=1,cols=2,width=Mm(182)); ht.autofit=False
ht.columns[0].width=Mm(47); ht.columns[1].width=Mm(135)
logo=ROOT/'apps/worker/assets/naabsa-logo.jpg'
if not logo.exists():
    logo=ROOT/'apps/web/public/naabsa-logo.jpg'
ht.cell(0,0).paragraphs[0].add_run().add_picture(str(logo),width=Mm(43))
p=ht.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for i,line in enumerate(['MARINE SURVEYORS & CONSULTANTS','Main Brazilian Ports']):
    r=p.add_run(('\n' if i else '')+line); r.font.name='Times New Roman'; r.font.size=Pt(13 if not i else 9); r.font.color.rgb=RGBColor.from_string('002060')
p=header.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=Pt(1); p.add_run().font.size=Pt(1)
prop(prop(ht._tbl.tblPr,'tblBorders'),'bottom',val='single',sz=4,color='002060')
fp=section.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.tab_stops.add_tab_stop(Mm(175),WD_TAB_ALIGNMENT.RIGHT)
fp.add_run('naabsa.com | surveyors@naabsa.com\t').font.size=Pt(8)
fld=prop(fp._p,'fldSimple',instr=' PAGE '); prop(prop(fld,'r'),'t').text='1'
prop(prop(fp._p.get_or_add_pPr(),'pBdr'),'top',val='single',sz=4,color='002060')

# Cover.
p=text('433 Ana Costa Avenue\tTelephone: +55 13 33940655\nSantos / Brazil\temail: surveyors@naabsa.com\n11060-003\twww.naabsa.com',size=9)
p.paragraph_format.tab_stops.add_tab_stop(Mm(112))
text('Survey Report','Title',center=True)
p=text('Ref: ',center=True); manual(p,'report_reference','Report reference')
text('Bunker Quantity Survey',center=True,size=12,bold=True)
text('MV {vessel_name}',center=True,size=20,bold=True)
text('{survey_location}',center=True,size=11)
text('{survey_date}',center=True,size=10)
photo_area('cover',True)
text('PERSON / COMPANIES CONTACTED',bold=True,size=11)
contacts=table(['',''],[54,128],borders=False)
contacts._tbl.remove(contacts.rows[0]._tr)
for label,name,val in [('Customer','customer',None),('Undersigned surveyor','surveyor','NAABSA Marine Surveyors\n{surveyor}'),('Suppliers / barge','supplier','{barge_name}'),("Vessel’s command",'master','Chief Engineer: {chief_engineer}')]:
    cells=row(contacts,[label,val or ''],size=9)
    if name!='surveyor':
        p=cells[1].paragraphs[0]
        if val: p.add_run('\n')
        manual(p,name)
mark(text(''),'bqs_cover_end')

p=text('Contents','Heading 1'); p.paragraph_format.page_break_before=True
for name,title,level in headings:
    p=text(title+'\t',size=9)
    field=prop(p._p,'fldSimple',instr=' PAGEREF bqs_'+name+' \\h ')
    rr=prop(field,'r'); prop(rr,'t').text='{toc_'+name+'}'
    p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.left_indent=Mm(6 if level==2 else 0)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(180),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),'bqs_'+name)
    for r in list(p._p):
        if r.tag!=qn('w:pPr'): link.append(r)
    p._p.append(link)

heading('s1',True)
text('A Bunker Quantity Survey was carried out on MV {vessel_name} at {survey_location} on {survey_date} to determine the quantity of {fuel_grade} received from barge {barge_name}.')
p=text('Appointing party / operational background: '); manual(p,'background')
heading('s2')
t=table(['',''],[46,136],borders=False); t._tbl.remove(t.rows[0]._tr)
for label,name in [('Flag','flag'),('Port of registry','registry'),('Call sign','call_sign'),('IMO number','imo'),('Type','vessel_type'),('Delivered','delivered'),('LOA (m)','loa')]:
    cells=row(t,[label,''],size=9); manual(cells[1].paragraphs[0],name,label)
heading('s3')
text('The surveyor boarded the vessel on {arrival_date} at {arrival_start} local time. The meeting with the Chief Engineer was recorded from {meeting_start} to {meeting_end}. Bunkering operations were recorded from {bunkering_start} to {bunkering_end} on {bunkering_date}.')
heading('s4'); heading('s4_1')
text('Opening measurements of the vessel’s tanks were recorded on {vessel_open_date}, from {vessel_open_start} to {vessel_open_end} local time.')
tank('vessel_open')
text('Logbook ({fuel_grade}): {vessel_logbook} MT\nNAABSA figures (IN AIR): {vessel_open_air} MT\nDifference (NAABSA − logbook): {vessel_logbook_difference} MT',size=9)

heading('s4_2',True)
text('Gross observed volumes are shown in the vessel’s tank measurement table. Trim correction applied: {vessel_open_trim_corrected}.')
photo_area('sounding_tables')
heading('s4_3'); text('{vessel_thermometer}.')
p=text('Engine room temperature / sea water temperature: '); manual(p,'ambient_temperatures')
heading('s4_4')
text('Densities recorded at 15°C in the opening vessel tank measurements: {vessel_densities}.')
heading('s4_5')
text('Opening measurements of barge {barge_name} were recorded on {barge_open_date}, from {barge_open_start} to {barge_open_end} local time.')
tank('barge_open')
text('Surveyor’s figures (IN AIR): {barge_open_air} MT\nFlowmeter reading: {barge_open_flowmeter}',size=9)
p=text('Barge inspector’s figures (MT): ',size=9); manual(p,'barge_inspector_open')
heading('s4_6'); text('Gross observed volumes are shown in the barge’s tank measurement table. Trim correction applied: {barge_open_trim_corrected}.')
heading('s4_7'); text('Thermometer: {barge_thermometer}.')
heading('s4_8')
t=table(['Temperature °C','Specific gravity as per BDN'],[45,90]); row(t,['15','{bdn_density_15}'],center=True); row(t,['20','{bdn_density_20}'],center=True)

heading('s5',True); heading('s5_1')
text('Closing measurements of the vessel’s tanks were recorded on {vessel_close_date}, from {vessel_close_start} to {vessel_close_end} local time.')
tank('vessel_close')
text('Total quantity received as per the vessel’s tanks (IN AIR):\nInitial: {vessel_open_air} MT    Final: {vessel_close_air} MT\nQuantity received: {vessel_received} MT',size=9)
heading('s5_2')
text('Closing measurements of barge {barge_name} were recorded on {barge_close_date}, from {barge_close_start} to {barge_close_end} local time.')
tank('barge_close')
text('Surveyor’s closing figures (IN AIR): {barge_close_air} MT\nQuantity delivered per barge measurements: {barge_delivered} MT\nFlowmeter reading: {barge_close_flowmeter}',size=9)
p=text('Barge inspector’s figures (MT): ',size=9); manual(p,'barge_inspector_close')

heading('s6',True)
t=table(['Description','Quantity (MT)'],[130,52],size=10)
for label,name in [('BDN','bdn_quantity'),('Surveyor — vessel tanks (IN AIR)','vessel_received'),('Surveyor — barge tanks (IN AIR)','barge_delivered'),('Difference — vessel less BDN','delivery_difference')]:
    row(t,[label,'{'+name+'}'],size=10)
text('Difference relative to BDN: {delivery_difference_pct}%.')
text('Vessel and barge quantities are reported separately. The vessel figure is the closing quantity less the opening quantity; the barge figure is the opening quantity less the closing quantity. Line condition: {line_condition}.')
p=text('Surveyor’s remarks / any subsequent measurements: '); manual(p,'final_remarks')
heading('s7'); heading('s7_1'); photo_area('vessel_tanks')
heading('s7_2'); photo_area('barge_tanks')
heading('s7_3'); photo_area('sampling')

heading('s8',True)
text('The tank measurements are reproduced in sections 4 and 5. The time log and sampling record are reproduced below.')
p=text('Additional documents / attachment references: '); manual(p,'attachments')
heading('s8_1')
t=table(['Operation','Date','Start','End'],[98,32,26,26],size=8)
row(t,['{#time_log}{event}','{date}','{start}','{end}{/time_log}'],size=8)
heading('s8_2')
text('Method: {sampling_method}. Tanks: {sampling_tanks}.',size=9)
t=table(['','Vessel','Barge'],[66,58,58],size=8)
row(t,['Sampling point','{sample_vessel_point}','{sample_barge_point}'])
row(t,['Quantity','{sample_vessel_quantity}','{sample_barge_quantity}'])
for label,name in [('Chief Engineer','chief'),('Supplier representative','supplier'),('Laboratory','laboratory'),('MARPOL','marpol'),('Surveyor','surveyor')]:
    row(t,[label,'{sample_vessel_'+name+'}','{sample_barge_'+name+'}'])
text('Chief Engineer: {chief_engineer}\nSurveyor: {surveyor}',size=9)

doc.core_properties.title='Bunker Quantity Survey'
doc.core_properties.author='NAABSA'
doc.core_properties.subject='BQS editable model reconstructed from client PDF'
out=ROOT/'templates/bqs.clean.docx'
doc.save(out)
print(out)
